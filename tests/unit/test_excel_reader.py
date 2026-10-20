import os

from cosmoslib.excel_reader import read_named_range

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


def test_sheet_lookup_is_case_insensitive():
    # utilities/manifests/bc/20260714_r18.yaml's BC_SURROGATES job differs in case from the
    # actual "BC_Surrogates" sheet in this workbook - PROC IMPORT resolves that on Windows,
    # so read_named_range must too.
    path = os.path.join(REPO_ROOT, "curation", "package18", "R18_BC_Surrogates.xlsx")
    df = read_named_range(path, "BC_SURROGATES$")
    assert len(df) > 0
    assert "bc_id" in df.columns


def test_blank_header_columns_are_dropped_not_kept_as_duplicates():
    # R18_BC_SDTM_SC.xlsx's "SDTM_SC" sheet has ~400 trailing blank-header columns from
    # openpyxl's used-range detection - these must not survive as duplicate "" columns,
    # which would break pd.concat/merge across sheets.
    path = os.path.join(REPO_ROOT, "curation", "package18", "R18_BC_SDTM_SC.xlsx")
    df = read_named_range(path, "SDTM_SC$")
    assert "" not in df.columns
    assert not df.columns.duplicated().any()
    assert "vlm_group_id" in df.columns


def _workbook_with_hyperlink_on_empty_cell(path):
    """A sheet whose B3 (and A5:A39) are EMPTY but still carry a hyperlink - what Excel leaves behind when a
    cell's content is deleted but not its hyperlink (see the excel_reader module docstring). openpyxl will not
    save a hyperlink on a valueless cell, so the cells are saved with a placeholder value and then emptied in
    the sheet XML."""
    import re
    import zipfile

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SDTM_TEST"
    sheet.append(["vlm_group_id", "dec_id", "codelist"])
    sheet.append(["G1", "C100", "C200"])
    sheet.append(["G2", "placeholder", "C300"])
    sheet["B3"].hyperlink = "https://evsexplore.semantics.cancer.gov/evsexplore/concept/ncit/C999"
    sheet.append(["G3", None, None])
    for row in range(5, 40):
        sheet[f"A{row}"].value = "x"
        sheet[f"A{row}"].hyperlink = f"https://example.org/{row}"
    workbook.save(path + ".tmp")

    with zipfile.ZipFile(path + ".tmp") as source, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                # empty B3 and A5:A39: keep the cell element, drop its stored value
                text = re.sub(r'<c r="(B3|A(?:[5-9]|[1-3][0-9]))"([^>]*?) t="(?:s|inlineStr)"([^>]*)>.*?</c>',
                              r'<c r="\1"\2\3/>', text, flags=re.S)
                data = text.encode("utf-8")
            target.writestr(item, data)
    os.remove(path + ".tmp")


def test_hyperlink_on_empty_cell_stays_blank(tmp_path):
    # Regression: openpyxl's normal load mode turns a hyperlink on an empty cell into a cell VALUE equal to
    # the link target, so blank curation cells read as NCIt explore URLs and the port disagreed with SAS
    # (which reads them as blank). The reader must return only stored cell values.
    from openpyxl import load_workbook

    path = str(tmp_path / "hyperlinks.xlsx")
    _workbook_with_hyperlink_on_empty_cell(path)

    # sanity check that the fixture really reproduces the hazard in openpyxl's normal mode
    hazard = load_workbook(path)["SDTM_TEST"]
    assert str(hazard["B3"].value).startswith("https://evsexplore")

    df = read_named_range(path, "SDTM_TEST$")
    # the emptied A5:A39 cells still exist as blank rows; like SAS, the reader leaves filtering to the converters
    assert list(df["vlm_group_id"].dropna()) == ["G1", "G2", "G3"]
    assert df["vlm_group_id"].iloc[3:].isna().all()
    assert df.loc[df["vlm_group_id"] == "G2", "dec_id"].isna().all()
    assert df.loc[df["vlm_group_id"] == "G1", "dec_id"].iloc[0] == "C100"
    assert not any(str(v).startswith("http") for column in ("dec_id", "codelist") for v in df[column].dropna())


def test_rows_shorter_than_header_are_padded(tmp_path):
    from openpyxl import Workbook

    path = str(tmp_path / "short_rows.xlsx")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "S"
    sheet.append(["a", "b", "c"])
    sheet.append(["1"])  # trailing cells never stored
    workbook.save(path)

    df = read_named_range(path, "S$")
    assert df.loc[0, "a"] == "1"
    assert df["b"].isna().all() and df["c"].isna().all()


def test_r18_brca_workbook_has_no_phantom_url_cells():
    # R18_BC_DSS_BrCa_TAUG.xlsx's SDTM_BrCa sheet ends at row 139 but carries ~5,500 leftover hyperlinks on
    # empty cells, e.g. J120 (AJCCV7_AJCC105 / RSTEST dec_id), which Excel and SAS read as blank.
    path = os.path.join(REPO_ROOT, "curation", "package18", "R18_BC_DSS_BrCa_TAUG.xlsx")
    df = read_named_range(path, "SDTM_BrCa$")

    assert len(df) < 200  # the real table; before the fix hyperlinks on empty cells extended it to 2,509 rows
    for column in ("bc_id", "dec_id", "codelist"):
        assert not any(str(v).startswith("http") for v in df[column].dropna()), column

    row = df[(df["vlm_group_id"] == "AJCCV7_AJCC105") & (df["sdtm_variable"] == "RSTEST")]
    assert len(row) == 1
    assert row["dec_id"].isna().all()
